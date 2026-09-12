"""
PAIMANA AI — Professional PowerPoint Presentation Generator (.pptx)
Builds an executive pitch deck from COMPLETE_PLATFORM_DOCUMENTATION.md
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette (MoSPI Slate & GovTech Modern)
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)     # Slate 700
    TEXT_WHITE = RGBColor(248, 250, 252)   # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    ACCENT_BLUE = RGBColor(56, 189, 248)   # Sky Blue
    ACCENT_GOLD = RGBColor(245, 158, 11)   # Amber/Gold
    ACCENT_GREEN = RGBColor(16, 185, 129)  # Emerald
    ACCENT_RED = RGBColor(239, 68, 68)     # Crimson Red

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="PAIMANA AI · SMART INDIA HACKATHON 26103"):
        # Category label
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE
        p_cat.font.name = "Segoe UI"

        # Slide Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Segoe UI"

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.45))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = ACCENT_GOLD
            p.font.name = "Segoe UI"
        return card

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_BLUE
    bar.line.fill.background()

    # Title box
    tb = slide1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "PAIMANA AI"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Segoe UI"

    p2 = tf.add_paragraph()
    p2.text = "AI-Powered Infrastructure Risk Monitoring & Predictive Intervention Platform"
    p2.font.size = Pt(22)
    p2.font.color.rgb = ACCENT_BLUE
    p2.font.name = "Segoe UI"

    # Meta card
    add_card(slide1, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.0), border_color=ACCENT_BLUE)
    tb_meta = slide1.shapes.add_textbox(Inches(1.1), Inches(4.7), Inches(11.1), Inches(1.6))
    tf_meta = tb_meta.text_frame
    tf_meta.word_wrap = True

    m1 = tf_meta.paragraphs[0]
    m1.text = "Ministry: Ministry of Statistics and Programme Implementation (MoSPI) · IPMD"
    m1.font.size = Pt(14)
    m1.font.bold = True
    m1.font.color.rgb = TEXT_WHITE

    m2 = tf_meta.add_paragraph()
    m2.text = "Smart India Hackathon (SIH 2026) · Problem Statement: SIH26103"
    m2.font.size = Pt(13)
    m2.font.color.rgb = ACCENT_GOLD

    m3 = tf_meta.add_paragraph()
    m3.text = "Core Tech: Gradient Boosted Trees (Model B) · Genuine SHAP TreeExplainer · Counterfactual What-If Simulator · Offline-First Field Telemetry"
    m3.font.size = Pt(12)
    m3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: THE PROBLEM & NEED
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "The National Challenge: Mega-Infrastructure Delays & Cost Overruns")

    # Card 1: Ground Reality
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), title="1. Ground Reality")
    tb2_1 = slide2.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(3.2), Inches(4.3))
    tf2_1 = tb2_1.text_frame
    tf2_1.word_wrap = True
    bullets1 = [
        ("Massive Scale: ", "MoSPI monitors Central Sector Projects (₹150 Cr+) across Highways, Railways, Bridges, and Power."),
        ("Chronic Delays: ", "Average schedule slippage of 30 to 45 months beyond sanctioned completion dates."),
        ("Cost Drift: ", "Cumulative cost overruns of 18%–24% across unmonitored infrastructure packages."),
    ]
    for i, (bold, txt) in enumerate(bullets1):
        p = tf2_1.paragraphs[0] if i == 0 else tf2_1.add_paragraph()
        run1 = p.add_run()
        run1.text = bold
        run1.font.bold = True
        run1.font.color.rgb = ACCENT_BLUE
        run1.font.size = Pt(12)
        run2 = p.add_run()
        run2.text = txt
        run2.font.color.rgb = TEXT_MUTED
        run2.font.size = Pt(12)
        p.space_after = Pt(10)

    # Card 2: Why Current Systems Fail
    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), title="2. The Root Bottlenecks", border_color=ACCENT_RED)
    tb2_2 = slide2.shapes.add_textbox(Inches(5.0), Inches(2.3), Inches(3.2), Inches(4.3))
    tf2_2 = tb2_2.text_frame
    tf2_2.word_wrap = True
    bullets2 = [
        ("Paper-Driven MPRs: ", "Monthly progress reports arrive 30+ days late — purely retrospective."),
        ("Disbursement Fraud: ", "Running Account (RA) bills paid out without verified ground cross-section audits."),
        ("Opaque Subcontracting: ", "Multi-tier petty sub-letting creates invisible labor and wage bottlenecks."),
        ("Unresolved Litigations: ", "Right-of-Way (RoW) disputes stall construction without prompt administrative notice."),
    ]
    for i, (bold, txt) in enumerate(bullets2):
        p = tf2_2.paragraphs[0] if i == 0 else tf2_2.add_paragraph()
        run1 = p.add_run()
        run1.text = bold
        run1.font.bold = True
        run1.font.color.rgb = ACCENT_RED
        run1.font.size = Pt(12)
        run2 = p.add_run()
        run2.text = txt
        run2.font.color.rgb = TEXT_MUTED
        run2.font.size = Pt(12)
        p.space_after = Pt(8)

    # Card 3: The PAIMANA Paradigm Shift
    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), title="3. The PAIMANA Solution", border_color=ACCENT_GREEN)
    tb2_3 = slide2.shapes.add_textbox(Inches(9.0), Inches(2.3), Inches(3.3), Inches(4.3))
    tf2_3 = tb2_3.text_frame
    tf2_3.word_wrap = True
    bullets3 = [
        ("Anticipatory ML: ", "Predicts delays 6 months before milestones slip using 12 friction signals."),
        ("Zero Black-Box Opacity: ", "100% mathematically reconciled SHAP attribution explaining every point."),
        ("Policy Simulation: ", "Counterfactual 'What-If' engine testing intervention ROI before spending capital."),
        ("Closed-Loop Telemetry: ", "Offline-first mobile logging directly from remote chainage site engineers."),
    ]
    for i, (bold, txt) in enumerate(bullets3):
        p = tf2_3.paragraphs[0] if i == 0 else tf2_3.add_paragraph()
        run1 = p.add_run()
        run1.text = bold
        run1.font.bold = True
        run1.font.color.rgb = ACCENT_GREEN
        run1.font.size = Pt(12)
        run2 = p.add_run()
        run2.text = txt
        run2.font.color.rgb = TEXT_MUTED
        run2.font.size = Pt(12)
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: SYSTEM ARCHITECTURE
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture: Decoupled, Robust & Enterprise-Ready")

    cols = [
        ("1. Ground Telemetry", Inches(0.8), "Field Officer Mobile PWA\n\n• Daily telemetry intake\n• Work Status (Running/Stalled/Off)\n• Material inventory & consumption\n• Geotagged site photo proof\n• Offline localStorage queue"),
        ("2. Cloud Persistence", Inches(3.9), "Supabase PostgreSQL 15\n\n• 6 Strongly-typed tables\n• Role-Based Access Control (RLS)\n• High-concurrency indices\n• Real-time database streams\n• Secure JWT Auth Handshake"),
        ("3. Predictive Engine", Inches(7.0), "Python FastAPI & Model B\n\n• XGBRegressor & XGBClassifier\n• SHAP TreeExplainer\n• Feature enrichment pipeline\n• Two-Tier dynamic project lookup\n• Batch write-back sync"),
        ("4. Executive Shell", Inches(10.1), "Vite + React 18 Web App\n\n• MoSPI Executive Dashboard\n• Priority Intervention Queue\n• Cross-sector Benchmarking\n• Clause 14.2 Billing Alerts\n• What-If Counterfactual Sandbox"),
    ]
    for title, left, desc in cols:
        add_card(slide3, left, Inches(1.6), Inches(2.8), Inches(5.2), title=title)
        tb = slide3.shapes.add_textbox(left + Inches(0.2), Inches(2.3), Inches(2.4), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = "Segoe UI"

    # ==========================================
    # SLIDE 4: MOSPI EXECUTIVE VIEW
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "MoSPI Executive View: Portfolio Governance & Real-Time Situational Awareness")

    # Left Card: Executive Dashboard
    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="1. Central Executive Dashboard (/dashboard)")
    tb4_l = slide4.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True
    features_l = [
        "Executive KPI Strip: Total monitored projects, High-Risk packages, Time overrun probability, and cumulative Cost Drift Exposure.",
        "Sector Breakdown: Instant filtering across Roads, Railways, Bridges, and Power corridors.",
        "Proportional Risk Bar: Dynamic distribution into High Risk (Red >= 66), Watch List (Amber 34-65), and On Track (Green <= 33).",
        "Priority Watch Table: Instant visibility into physical milestone deficits vs financial burn rate.",
        "One-Click Drilldown: Direct navigation to granular project diagnostics.",
    ]
    for i, item in enumerate(features_l):
        p = tf4_l.paragraphs[0] if i == 0 else tf4_l.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # Right Card: Project Diagnostics & Priority Queue
    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="2. Project Diagnostics (/projects/:id)")
    tb4_r = slide4.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True
    features_r = [
        "Radial Risk Gauge: Instant visual gauge showing composite risk score (0-100).",
        "Days Flagged Counter: Tracks persistent bottleneck severity over reporting cycles.",
        "6-Month Historical Risk Curve: Interactive trend chart showing risk velocity.",
        "Automated Policy Recommendations: Generates legally actionable contract interventions under FIDIC/MoRTH Clause 14.2.",
        "Priority Queue (/priority-queue): Automated urgency ranking for Public Investment Board (PIB) escalations.",
    ]
    for i, item in enumerate(features_r):
        p = tf4_r.paragraphs[0] if i == 0 else tf4_r.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: GENUINE SHAP EXPLAINABILITY
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Genuine SHAP Explainability Engine: No Black-Box Opacity")

    # Top banner: The Guarantee
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.4), title="Mathematical Reconciliation Guarantee", border_color=ACCENT_BLUE)
    tb5_t = slide5.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.8))
    tf5_t = tb5_t.text_frame
    p_eq = tf5_t.paragraphs[0]
    p_eq.text = "base_value + Σ SHAP_i = raw_prediction   (Exact to 10^-5 floating precision)"
    p_eq.font.size = Pt(18)
    p_eq.font.bold = True
    p_eq.font.color.rgb = ACCENT_GOLD
    p_sub = tf5_t.add_paragraph()
    p_sub.text = "Every displayed factor traces directly to a real explainer.shap_values() calculation on trained XGBoost decision trees. Zero hardcoded rules."
    p_sub.font.size = Pt(11.5)
    p_sub.font.color.rgb = TEXT_MUTED

    # Bottom Left: Dual-Mode Concept
    add_card(slide5, Inches(0.8), Inches(3.2), Inches(5.6), Inches(3.6), title="Dual-Mode Attribution Engine")
    tb5_bl = slide5.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(5.2), Inches(2.8))
    tf5_bl = tb5_bl.text_frame
    tf5_bl.word_wrap = True
    items5_bl = [
        ("Risk Escalation Mode (pred >= base_value): ", "Ranks positive SHAP features driving troubled projects (e.g. NH-4471, Risk 95) toward failure."),
        ("Protective Mode (pred < base_value): ", "Ranks negative SHAP features pulling stable projects (e.g. RW-8802, Risk 18) below population baseline."),
        ("Executive Value: ", "Provides actionable best practices that can be audited and transferred across corridors."),
    ]
    for i, (bold, txt) in enumerate(items5_bl):
        p = tf5_bl.paragraphs[0] if i == 0 else tf5_bl.add_paragraph()
        r1 = p.add_run()
        r1.text = bold
        r1.font.bold = True
        r1.font.color.rgb = ACCENT_BLUE
        r1.font.size = Pt(11.5)
        r2 = p.add_run()
        r2.text = txt
        r2.font.color.rgb = TEXT_MUTED
        r2.font.size = Pt(11.5)
        p.space_after = Pt(6)

    # Bottom Right: Top Feature Drivers
    add_card(slide5, Inches(6.8), Inches(3.2), Inches(5.7), Inches(3.6), title="Key Friction Signals Identified by SHAP")
    tb5_br = slide5.shapes.add_textbox(Inches(7.0), Inches(3.8), Inches(5.3), Inches(2.8))
    tf5_br = tb5_br.text_frame
    tf5_br.word_wrap = True
    items5_br = [
        "1. Billing Progress Mismatch (%): Financial claims outpacing physical work (+10.8 pts).",
        "2. Land Clearance Disputes: Right-of-Way legal litigation along chainage (+9.2 pts).",
        "3. Payment Delay Days: Disbursement lag straining contractor working capital.",
        "4. Subcontracting Depth: Excessive multi-tier petty sub-letting fragmentation.",
        "5. Design Scope Revisions: Repeated alignment and structural redesign orders.",
    ]
    for i, item in enumerate(items5_br):
        p = tf5_br.paragraphs[0] if i == 0 else tf5_br.add_paragraph()
        p.text = item
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(4)

    # ==========================================
    # SLIDE 6: WHAT-IF SIMULATOR
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "The What-If Simulator: Counterfactual Policy Sandbox")

    # Left: How it works
    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="1. Counterfactual Simulation Workflow")
    tb6_l = slide6.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf6_l = tb6_l.text_frame
    tf6_l.word_wrap = True
    steps = [
        "Step 1: Baseline Load — Retrieves project's real ground-research vector from database.",
        "Step 2: Baseline Inference — Runs Model B + TreeExplainer on unmodified project.",
        "Step 3: User Intervention — Official adjusts specific friction parameters (e.g. resolve land dispute, clear billing mismatch).",
        "Step 4: Simulated Inference — Runs Model B + TreeExplainer on counterfactual project.",
        "Step 5: Quantified Delta — System returns exact Risk Score Delta (ΔRisk) and updated SHAP attribution distribution.",
    ]
    for i, step in enumerate(steps):
        p = tf6_l.paragraphs[0] if i == 0 else tf6_l.add_paragraph()
        p.text = step
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right: Live Case Example
    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="2. Case Study: NH-4471 Highway Intervention", border_color=ACCENT_GREEN)
    tb6_r = slide6.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf6_r = tb6_r.text_frame
    tf6_r.word_wrap = True

    p = tf6_r.paragraphs[0]
    p.text = "Baseline Assessment:"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = ACCENT_RED

    p = tf6_r.add_paragraph()
    p.text = "• Risk Score: 95 / 100  |  Status: High Critical\n• Top Driver: Billing Mismatch (37.5%) + RoW Dispute\n• Time Overrun Probability: 99.6%"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(12)

    p = tf6_r.add_paragraph()
    p.text = "Simulated Intervention:"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = ACCENT_BLUE

    p = tf6_r.add_paragraph()
    p.text = "• Intervene on: Payment delay cut to 10 days\n• Intervene on: Land Clearance set to 'Clear'\n• Intervene on: Billing mismatch audited down to 2%"
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = Pt(12)

    p = tf6_r.add_paragraph()
    p.text = "Counterfactual Result: Risk drops 95 -> 71 pts (Improved!)\nAdministration validates exactly which intervention delivers maximum risk reduction before spending public money."
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    # ==========================================
    # SLIDE 7: FINANCIAL FRAUD & BILLING ALERTS
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Billing Alerts & Fraud Prevention: Enforcing Clause 14.2")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), title="Financial Disbursement vs Ground Progress Verification")
    tb7 = slide7.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.4))
    tf7 = tb7.text_frame
    tf7.word_wrap = True

    b_items = [
        ("The Core Financial Risk: ", "Concessionaires submitting inflated Running Account (RA) bills claiming milestone completion before third-party physical cross-section verification."),
        ("Billing Progress Mismatch Formula: ", "Mismatch % = [(Total Claimed - Expected Physical) / Expected Physical] * 100"),
        ("Automated Trigger 1 (>10% Mismatch): ", "System issues an amber financial warning to the Project Director."),
        ("Automated Trigger 2 (>25% Mismatch): ", "System automatically flags the project in MoSPI dashboard, applies Clause 14.2 disbursement freeze recommendation, and forces an emergency on-site joint re-measurement audit."),
        ("Impact on ML Risk Engine: ", "Billing progress mismatch directly feeds Model B feature vectors, instantly raising risk score and triggering SHAP factor escalation."),
    ]
    for i, (bold, txt) in enumerate(b_items):
        p = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        r1 = p.add_run()
        r1.text = bold
        r1.font.bold = True
        r1.font.color.rgb = ACCENT_GOLD
        r1.font.size = Pt(13)
        r2 = p.add_run()
        r2.text = txt
        r2.font.color.rgb = TEXT_WHITE
        r2.font.size = Pt(13)
        p.space_after = Pt(12)

    # ==========================================
    # SLIDE 8: FIELD OFFICER VIEW
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Field Officer View: Ground Telemetry & Mobile Verification")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="1. On-Site Inspection Intake (/field-entry)")
    tb8_l = slide8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf8_l = tb8_l.text_frame
    tf8_l.word_wrap = True
    f_l = [
        "Work Status Logging: One-tap toggle for Running (Green), Stalled (Red), or Off (Gray).",
        "Standardized Delay Taxonomy: Land Dispute, Monsoon Flooding, Material Shortage, Contractor Cashflow, Subcontractor Wage Strike, or Design Change.",
        "Material Stockpile Tracking: Daily logs of cement bags, structural steel tonnage, and coarse aggregate on site.",
        "Geotagged Photo Proof: Attaches camera photo with exact chainage marker (e.g. Km 44.200 - Pier 12) and inspection timestamp.",
        "Today's Tasks (/field-tasks): Action checklist for mandatory joint measurements and slump tests.",
    ]
    for i, item in enumerate(f_l):
        p = tf8_l.paragraphs[0] if i == 0 else tf8_l.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="2. Offline-First Resilience & Sync", border_color=ACCENT_BLUE)
    tb8_r = slide8.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf8_r = tb8_r.text_frame
    tf8_r.word_wrap = True
    f_r = [
        "Zero Connectivity Blackout Protection: Remote mountain passes and rail cuts lose cellular connectivity regularly.",
        "Local Storage Engine: When offline, submissions are instantly serialized into device localStorage (paimana_submissions).",
        "Visual Offline Banner: Alerts the site officer that telemetry is safely stored locally on device.",
        "Automatic Cloud Sync: Device automatically flushes queued submissions to Supabase PostgreSQL as soon as network reconnects.",
        "Compliance Grading: Dynamically calculates inspection streak and letter grade (A+, A, B, C, Critical).",
    ]
    for i, item in enumerate(f_r):
        p = tf8_r.paragraphs[0] if i == 0 else tf8_r.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 9: DATA CALIBRATION & TRANSPARENCY
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Data Governance & Rigorous Synthetic Calibration")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), title="Ethical AI & MoSPI-Aggregate Calibration Disclosure")
    tb9 = slide9.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.4))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    c_items = [
        ("Classified Data Boundary: ", "Real row-level PAIMANA government infrastructure project telemetry and dispute documentation are restricted under Government of India data governance protocols."),
        ("5,000-Project Calibrated Dataset: ", "Model B was trained on an empirical synthetic dataset generated by ml/generate_dataset.py, calibrated strictly to MoSPI's published quarterly aggregate statistics:"),
        ("  • Capital Distribution: ", "Power-law and log-normal distributions mirroring central sector packages (>= ₹150 Cr)."),
        ("  • Historical Delay Rates: ", "Calibrated to historical MoSPI flash report averages (~44% delay incidence)."),
        ("  • Sector Proportions: ", "Highways, Railways, Bridges, and Power sampled proportional to national pipeline budgets."),
        ("Model Card Documentation: ", "Full disclosure on model performance, feature distributions, and ethical limitations documented in model_card.md."),
        ("Two-Tier Resolution: ", "When live projects are created in Supabase, the engine dynamically ingests their live telemetry without requiring model retraining."),
    ]
    for i, (bold, txt) in enumerate(c_items):
        p = tf9.paragraphs[0] if i == 0 else tf9.add_paragraph()
        r1 = p.add_run()
        r1.text = bold
        r1.font.bold = True
        r1.font.color.rgb = ACCENT_GOLD
        r1.font.size = Pt(12.5)
        r2 = p.add_run()
        r2.text = txt
        r2.font.color.rgb = TEXT_WHITE
        r2.font.size = Pt(12.5)
        p.space_after = Pt(6)

    # ==========================================
    # SLIDE 10: CONCLUSION & NATIONAL IMPACT
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Conclusion: Transforming National Infrastructure Delivery")

    # 3 Summary impact columns
    c1 = add_card(slide10, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), title="Public Capital Safeguard", border_color=ACCENT_BLUE)
    tb10_1 = slide10.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(3.2), Inches(4.3))
    tf10_1 = tb10_1.text_frame
    tf10_1.word_wrap = True
    p = tf10_1.paragraphs[0]
    p.text = "• Eliminates premature RA bill disbursement leakage.\n\n• Enforces Clause 14.2 financial discipline.\n\n• Saves hundreds of crores in unnecessary cost escalation claims."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    c2 = add_card(slide10, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), title="Anticipatory Action", border_color=ACCENT_GOLD)
    tb10_2 = slide10.shapes.add_textbox(Inches(5.0), Inches(2.3), Inches(3.2), Inches(4.3))
    tf10_2 = tb10_2.text_frame
    tf10_2.word_wrap = True
    p = tf10_2.paragraphs[0]
    p.text = "• Moves MoSPI from 30-day lag to real-time predictive risk.\n\n• Identifies Right-of-Way and utility disputes months before work halts.\n\n• Recommends specific administrative interventions."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    c3 = add_card(slide10, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2), title="Ready for Scale", border_color=ACCENT_GREEN)
    tb10_3 = slide10.shapes.add_textbox(Inches(9.0), Inches(2.3), Inches(3.3), Inches(4.3))
    tf10_3 = tb10_3.text_frame
    tf10_3.word_wrap = True
    p = tf10_3.paragraphs[0]
    p.text = "• Production-ready Supabase PostgreSQL backend.\n\n• Offline-first mobile field logging.\n\n• Scalable FastAPI ML inference service.\n\n• Ready for immediate pilot deployment across IPMD central projects."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    # Save presentation
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PAIMANA_AI_SIH26103_Presentation.pptx")
    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated at: {output_path}")
    return output_path

if __name__ == "__main__":
    create_deck()
